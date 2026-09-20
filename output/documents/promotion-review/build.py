from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path
D=Document(); sec=D.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(1.8); sec.left_margin=sec.right_margin=Cm(1.9)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 s=D.styles[name]; s.font.name='Arial Unicode MS'; s.font.color.rgb=RGBColor(0,0,0); s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
 s.font.size=Pt(10.5 if name=='Normal' else 21 if name=='Title' else 16 if name=='Heading 1' else 12)
 s.paragraph_format.space_after=Pt(7); s.paragraph_format.line_spacing=1.15
for style in D.styles:
 for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
D.styles['Normal'].paragraph_format.widow_control=True
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
f.add_run('Tenant A 产品确认稿  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
for r in f.runs:r.font.size=Pt(9)
D.core_properties.title='Tenant A 推广规则产品确认稿';D.core_properties.author='';D.core_properties.subject='配置快照 计算案例 特殊情况与产品确认'
def p(t):return D.add_paragraph(t)
def h(t):D.add_heading(t,2)
def page(t):
 q=D.add_heading(t,1);q.paragraph_format.page_break_before=True
def table(headers,rows,widths):
 t=D.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for i,a in enumerate(headers):t.rows[0].cells[i].text=a
 for row in rows:
  for i,a in enumerate(row):t.add_row() if False else None
  cells=t.add_row().cells
  for i,a in enumerate(row):cells[i].text=str(a)
 for n,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');pr.append(ns)
  if n==0:pr.append(OxmlElement('w:tblHeader'))
  for i,c in enumerate(row.cells):
   c.width=Cm(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tc=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF4' if n==0 else 'FFFFFF');tc.append(sh)
   mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
   tc.append(mar)
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1.08
    if len(headers)>3:par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in par.runs:r.font.size=Pt(9.5);r.bold=n==0
 pr=t._tbl.tblPr; borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 pr.append(borders)
 D.add_paragraph().paragraph_format.space_after=Pt(0)
D.add_heading('Tenant A 推广规则产品确认稿',0)
p('年费返还  年费佣金  激活佣金')
p('请产品部核对本文所列当前行为，并在末尾确认清单填写结果与修改意见。本文记录配置及代码的现行逻辑，不代表产品已批准，也不代表已完成全部特殊场景的资金验收。')
p('配置读取时间：2026年9月19日 17:42（UTC+8）\n适用公司：Tenant A　公司时区：Asia/Kuala_Lumpur\n依据环境：本地 card_mock 数据库及当前代码　文档版本：v1.0')
D.add_heading('一 当前配置',1)
table(['等级','年费','年费佣金','激活标准','返费门槛'],[
['普通会员','无年费','0%','直推 20','不适用'],['万事达1级','1,000','30%','50','100'],['万事达2级','2,000','40%','60','200'],['万事达3级','5,000','50%','70','500'],['万事达4级','10,000','60%','80','1,000'],['万事达5级','20,000','70%','90','2,000'],['万事达6级','50,000','80%','100','5,000'],['万事达7级','100,000','90%','110','10,000'],['万事达8级','200,000','100%','120','20,000']],[3.4,3.2,3.4,3.4,3.8])
p('金额单位均为 USDT；八档代理均已启用。普通会员保证金为 300，退款等待期为 30 天，保证金不是年费。激活标准用于推荐链级差计算，不是每个上级均可获得全额。')
p('返年费进度按直推 1、间推 0.5 折算。已购周期采用保存的购买条款；修改后台配置不会直接重写既有周期。读取时仅有 1 个有效一级代理周期，其条款与本表一致。')
page('二 触发条件总览')
p('首次激活：同一公司内，账号第一次成功缴纳保证金或购买代理，只记录一次。充值钱包、注册和实名认证本身不构成激活。')
table(['用户行为','激活佣金','年费佣金','返费人数'],[
['注册或实名认证','无','无','不计'],['仅充值钱包','无','无','不计'],['首次缴保证金激活','有','无','计一次'],['未缴保证金直接首次买代理','无','钱包实付计佣','计一次'],['先缴保证金再买代理','不重复发','钱包实付计佣','不重复计'],['代理升级','无','本次钱包实付计佣','不计'],['代理到期后续费','无','本次钱包实付计佣','不计'],['保证金退款后重新缴纳','无','无','不计'],['保证金要求提高后补缴','无','无','不计'],['先买代理 过期后首次缴保证金','无','无','不计']],[6.6,2.8,4.2,3.6])
h('三种权益的区别')
p('有首次激活人数，不一定有激活佣金；有年费佣金，也不一定增加返年费人数。年费佣金由购买者向上级产生；年费返还是购买者本人达标后取回本周期未返年费。')
h('计数范围')
p('直属推荐人计 1；第二层及更远的所有上级均计 0.5，但仅计入各上级当时有效的付费周期。没有有效代理周期的上级，不因此获得返年费权益。关系按首次激活时的推荐链保存。')
h('首次无奖励也不会重算')
p('首次激活时没有上级、没有可得级差或相关标准为零，仍保留首次激活事实。后续重新缴纳、上级升级，不重新发激活佣金，也不重复计数。')
page('三 激活佣金规则与案例')
h('计算规则')
p('仅首次通过保证金激活触发激活佣金。沿推荐链由近到远，每级实际奖励为“本级标准减去前面已覆盖的最高标准”，小于零时按零计算。直属普通会员标准为 20；间推普通会员为零；有效代理使用其周期奖励标准。')
p('不扣减被激活用户自己的等级标准。同级或低级上级不重复拿钱，也不会增加已覆盖标准。以下链路从直属推荐人开始排列。')
table(['上级链路','实际到账','说明'],[
['普通 → 一级 → 三级','20、30、20','20；50−20；70−50'],['一级 → 一级 → 三级','50、0、20','同级不重复；三级补差'],['三级 → 一级 → 八级','70、0、50','低级不重复；120−70'],['普通 → 普通 → 二级','20、0、40','间推普通无奖励；60−20'],['八级直属','120','按直属八级标准']],[6,4,7.2])
h('退款与重新缴纳')
p('下级后来退保证金，不追回已发出的激活佣金，不撤销首次激活计数。退款后重新缴纳或要求提高后补缴，均不再产生激活奖励。')
h('代理到期与先后顺序')
p('上级代理到期后，无有效付费周期时按普通规则处理：直属可能获得 20，间推不获得。有效代理不能再缴保证金；先买代理激活的人，过期后即使第一次缴保证金，也不是第一次账号激活。')
h('产品应确认的经济差异')
p('假设直属上级为一级：下级先缴 300 保证金，上级得激活佣金 50；再买一级时抵扣 300，钱包付 700，年费佣金为 210，总佣金 260。若下级直接买一级，钱包付 1,000，上级获得 300 年费佣金，无激活佣金。两条路径下，购买者后续可返年费均为 1,000。')
page('四 年费佣金规则与案例')
h('计佣基数')
p('计佣基数仅为本次钱包实际支付的年费。保证金自动按“当前保证金与本次应付的较小值”抵扣；抵扣部分不计佣金。超过应付的保证金保留。')
table(['场景','应付总额','保证金抵扣','计佣基数'],[['直接买一级','1,000','0','1,000'],['已有300保证金 买一级','1,000','300','700'],['一级升二级 无剩余保证金','1,000','0','1,000'],['全部用保证金抵扣','按实际应付','全额','0']],[7,3.4,3.4,3.4])
h('分配资格与级差')
p('直属有效代理拿本人完整比例，不要求等级高于购买人；直属普通会员不拿年费佣金。间推代理必须不低于购买人本次购买后的等级，才可拿尚未覆盖的正比例差。不合格的中间层不占比例，购买人的比例也不参与扣减。')
table(['购买行为','上级链路','各级到账'],[['买一级 实付1,000','三级 → 六级','500、300'],['买三级 实付5,000','一级 → 二级 → 六级','1,500、0、2,500'],['买一级 实付1,000','五级 → 三级 → 八级','700、0、300'],['买一级 实付1,000','普通 → 二级 → 四级','0、400、200']],[5.7,6,5.5])
p('第二例：直属一级仍拿 30%；间推二级低于购买人三级，不参与；六级拿 80%−30%=50%。以上均假设无保证金抵扣且上级代理周期有效。')
h('升级 续费与返款')
p('升级按本次补差的钱包实付计佣，间推资格按升级后的等级判断。续费再次计佣，但不重复计首次激活人数。上级后来升级不补发历史差额；下级年费返还不追回上级佣金。每笔年费佣金保留 8 位小数，向下截取。')
page('五 自动返还年费')
p('进度 = 直推首次激活人数 + 所有间推首次激活人数 × 0.5。可返金额 = 本周期累计已付年费总额 − 本周期已返金额。已付总额包含保证金抵扣部分。达标后自动返入 USDT 钱包，无需手动申请或审核。')
table(['一级达标组合','直推','间推','折算进度'],[['全直推','100','0','100'],['全间推','0','200','100'],['混合','60','80','100']],[6.2,3.6,3.6,3.8])
h('时间边界')
p('仅计本人付费周期开始时刻及之后、到期时刻之前发生的首次激活。周期前已激活不计；早注册但本周期才首次激活可以计。恰好到期时发生的激活不计入旧周期。新周期不重复使用老团队已激活的人数。')
h('升级与续费案例')
table(['情况','当前处理'],[['一级未返 升二级','保留进度与到期日，门槛变200；达标可返累计2,000。'],['一级已返1,000 再升二级','仍补差1,000；达到200后只再返1,000。'],['升级时已经满足新门槛','购买完成后检查本人进度，可立即触发未返差额返还。'],['到期前一天升级','付完整档位差价，不按剩余天数折算，不延长有效期。'],['未到期想续费或降级','当前只允许升更高等级，不允许同级提前续费或降级。'],['到期后购买','按当时配置付费，开启新一年；周期内首次激活重新累计。']],[5.2,12])
p('返还年费不取消代理资格，等级仍在原有效期内有效。升级差价按已购买档位计算，不因之前已返年费而改收全额。返费待处理期间暂不允许升级。新一年按公司时区计算，遇日期溢出按不溢出的周年日期处理。')
page('六 到账与异常处理')
h('自动返还失败')
p('首次激活交易内先保存达标返还记录，交易提交后尝试到账。失败保留待处理权益，定时任务配置为每分钟重试。已捕获权益即使后来到期也可继续补发；同一权益重复处理不重复返款。')
p('恢复任务仅处理已有待处理记录，不是扫描所有历史周期补建遗漏权益的全量补算。任务配置存在不等于部署环境调度始终正常，本次未验证实际调度运行。')
h('佣金与钱包')
p('两类佣金直接进入 USDT 可用余额，无需再次转入钱包。累计佣金仅用于统计，不再额外计入总资产。年费返还单独统计，不算佣金收入。没有 USDT 钱包时，佣金收取可自动创建钱包，无需先通过 KYC；支出仍受正常资格限制。')
p('推荐人自己申请退保证金、取消退款或退款完成，不单独阻止佣金收入。钱包停用不会自动恢复，但在底层账本账户仍正常时，可以继续收佣；钱包状态与底层账本账户状态必须区分。')
h('资金事务边界')
p('下级缴保证金或购买代理，与上级佣金分配处于同一事务。若收款账本账户不可用或佣金记账失败，下级付款可能整笔回滚，不能描述成“跳过异常上级继续付款”。自动年费返还则在源交易提交后结算，失败不会回滚已成功的激活。')
h('报价与配置变更')
p('年费报价有效期为 5 分钟。确认时重新检查等级版本、当前周期、到期状态、保证金余额及退款状态。发生变化需要重新报价；保证金退款或恢复处理中不能直接抵扣。已完成订单重复确认不重复扣款。')
h('依据与适用限制')
p('配置来源：本地 card_mock 的 Tenant A，公司 ID 为 01a09996-8c36-7288-bf98-889103088ba6；paid_promotion_levels、tenant_business_settings、有效 paid_promotion_cycles。配置快照时间见第一页。')
p('代码依据：PaidPromotionRewards（佣金资格及级差）、PaidPromotionPurchase（购买与升级）、PaidPromotionRebate（计数与返还）、RecordAccountActivation（首次事实）、CommissionAccounts（收佣）、RecoverPromotion（待处理补发）。相关规则说明为 UNIFIED_ACCOUNT_ACTIVATION、DIRECT_COMMISSION_RECEIPTS。')
p('本文用于确认产品规则，不等同于线上实际配置或全场景资金测试结论。本文不更改后台配置、用户权益、历史订单或账本。')
in_six=False
for para in D.paragraphs:
 if para.text=='六 到账与异常处理': in_six=True
 if in_six and para.style.name=='Normal':
  para.paragraph_format.line_spacing=1.05
  para.paragraph_format.space_after=Pt(5)
checks=[
('间推范围','第二层及更深层均计0.5，佣金按资格和级差向全部祖先分配。'),
('代理首次购买也计激活','直接买代理计首次激活人数，不发激活佣金，只按钱包实付发年费佣金。'),
('两条购买路径佣金不同','直属一级时，先缴300再买一级合计佣金260；直接买一级佣金300。'),
('退款与返费不追回佣金','保证金退款不撤销计数及激活佣金；年费返还不追回年费佣金。'),
('直属与间推资格不同','直属低等级仍拿本人完整比例；间推等级须不低于购买者本次等级。'),
('升级的时间与门槛','不延长有效期、不按剩余天数折价，保留人数进度但提高门槛。'),
('新周期需要新激活','老团队重复充值、缴保证金、续费不能为新周期重复计数。'),
('八级100%的成本模型','钱包实付最高可全部分给推荐链；购买者达标后还可返年费，包括抵扣保证金。'),
('补发的实际覆盖范围','只恢复已保存待处理权益，不自动全量补建历史遗漏权益。'),
('上级异常对下级影响','上级账本账户不可用导致佣金记账失败，可能使下级付款整笔回滚。')]
for part in range(2):
 page('七 产品逐项确认' if part==0 else '七 产品逐项确认续页')
 p('以下“当前行为”为已核对内容。请填写确认结果（认可／需调整）及修改意见；空白表示尚未确认。')
 for i,(title,behavior) in enumerate(checks[part*5:part*5+5],part*5+1):
  h(str(i)+' '+title);p('当前行为：'+behavior)
  p('确认结果：________________    修改意见：________________________\n________________________________________________________')
 if part==1:p('产品确认人：________________　确认日期：________________\n结论：______________________　后续事项：________________')
out=Path('output/documents/Tenant A 推广规则产品确认稿.docx');D.save(out);print(out.resolve())
